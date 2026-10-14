"""Importação de extratos. Os ficheiros são gerados aqui com dados fictícios,
com a mesma estrutura dos exportados pelos bancos (nunca extratos reais no git)."""
import io
import re
from datetime import datetime

import openpyxl
import pytest
import xlwt
from fastapi.testclient import TestClient

from app.parsers import ler_extrato
from tests.test_app import PW, _csrf, _post  # noqa: F401
from app.main import create_app


def xls_santander(linhas) -> bytes:
    """linhas: (data 'dd-mm-aaaa', descrição, montante, saldo), mais recente primeiro."""
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Saldos e movimentos")
    for i, t in enumerate(["Quarta-feira, 14 de Outubro de 2026", "TITULAR FICTICIO", "", "Saldos e movimentos",
                           "", "Listagem de Movimentos"]):
        ws.write(i, 0, t)
    for j, t in enumerate(["Data Operação", "Data valor", "Descrição", "Montante( EUR )",
                           "Saldo Contabilístico( EUR )"]):
        ws.write(6, j, t)
    for i, (d, desc, mont, saldo) in enumerate(linhas, start=7):
        for j, v in enumerate([d, d, desc, mont, saldo]):
            ws.write(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_activobank(linhas) -> bytes:
    """linhas: (datetime, descrição, valor, saldo), mais antigo primeiro."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in [("HISTÓRICO DE CONTA NÚMERO 11122233344", None), ("Moeda:", "EUR"), ("", ""), ("Tipo:", "Todos"),
              ("Data de:", datetime(2026, 10, 4)), ("Data até:", datetime(2026, 10, 14)), (None, None)]:
        ws.append(list(r))
    ws.append(["Data Lanc.", "Data Valor", "Descrição", "Valor", "Saldo"])
    for d, desc, v, s in linhas:
        ws.append([d, d, desc, v, s])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SANT = [  # duas compras iguais no mesmo dia: só o saldo as distingue
    ("13-10-2026", "Supermercado X", -26.2, 1000.00),
    ("12-10-2026", "Cafe Central", -1.5, 1026.20),
    ("12-10-2026", "Cafe Central", -1.5, 1027.70),
    ("10-10-2026", "Vencimento", 1500.0, 1029.20),
]


def test_parser_santander_ordem_cronologica_e_saldos():
    e = ler_extrato("descarga.xls", xls_santander(SANT))
    assert e.parser == "santander" and e.moeda == "EUR" and not e.avisos
    assert [m.montante for m in e.movimentos] == [150000, -150, -150, -2620]
    assert e.movimentos[-1].saldo == 100000


def test_parser_activobank():
    e = ler_extrato("mov.xlsx", xlsx_activobank([
        (datetime(2026, 10, 5), "COMPRA 1234 CP CP", -0.5, 999.5),
        (datetime(2026, 10, 6), "TRF MB WAY P/ FULANO", -95, 904.5),
    ]))
    assert e.parser == "activobank" and e.conta_ref == "11122233344" and not e.avisos
    assert e.movimentos[0].descricao == "COMPRA 1234 CP CP"


def test_saldo_inconsistente_gera_aviso():
    linhas = list(SANT)
    linhas[0] = ("13-10-2026", "Supermercado X", -26.2, 999.99)
    assert ler_extrato("d.xls", xls_santander(linhas)).avisos


@pytest.fixture
def admin(tmp_path):
    app = create_app(tmp_path, backups_automaticos=False)
    c = TestClient(app)
    _post(c, "/setup", nome="Admin", login="admin", password=PW, password2=PW)
    _post(c, "/contas/nova?entidade=1", nome="Santander ordem", tipo="ORDEM", moeda="EUR", instituicao_id="1")
    return c


def _importar(c, nome, conteudo, conta_id="1"):
    r = c.get("/importar")
    previa = c.post("/importar", data={"csrf": _csrf(r), "conta_id": conta_id},
                    files={"ficheiro": (nome, conteudo)})
    assert previa.status_code == 200, previa.text
    sha = re.search(r'name="sha" value="([0-9a-f]+)"', previa.text)
    return previa, (c.post("/importar/confirmar", data={"csrf": _csrf(previa), "conta_id": conta_id,
                                                        "sha": sha.group(1), "nome": nome},
                           follow_redirects=False) if sha else None)


def test_importar_dedup_e_saldo_inicial(admin):
    previa, r = _importar(admin, "descarga.xls", xls_santander(SANT))
    assert "Importar 4 movimentos" in previa.text
    assert r.status_code == 303
    pagina = admin.get("/contas/1/movimentos").text
    assert "Saldo inicial calculado" in pagina
    assert "1.000,00 €" in pagina            # saldo atual = saldo do banco
    assert "-470,80 €" in pagina             # saldo inicial = 1029,20 - 1500
    assert "difere do calculado" not in pagina

    # extrato com um dia a mais e sobreposição total: só o novo entra
    mais = [("14-10-2026", "Farmacia", -10.0, 990.0)] + SANT
    previa, r = _importar(admin, "descarga2.xls", xls_santander(mais))
    assert "Importar 1 movimentos" in previa.text and "já existe" in previa.text
    assert "Importados 1 movimentos novos (4 já existiam)" in admin.get(r.headers["location"]).text

    # anular a 2.ª importação remove só esse movimento
    hist = admin.get("/importar")
    admin.post("/importar/2/anular", data={"csrf": _csrf(hist)})
    assert "990,00 €" not in admin.get("/contas/1/movimentos").text


def test_banco_errado_avisa(admin):
    previa, _ = _importar(admin, "mov.xlsx", xlsx_activobank([(datetime(2026, 10, 5), "X", -1, 9)]))
    assert "parece ser do ActivoBank" in previa.text


def test_ficheiro_invalido(admin):
    r = admin.get("/importar")
    r = admin.post("/importar", data={"csrf": _csrf(r), "conta_id": "1"},
                   files={"ficheiro": ("x.csv", b"a;b\n1;2\n")})
    assert r.status_code == 400 and "não reconhecido" in r.text


def test_mesmo_saldo_repetido_no_mesmo_dia(admin):
    """Transferência, devolução e nova transferência igual: a 1.ª e a 3.ª linhas
    têm data, montante, descrição e saldo iguais, mas são movimentos distintos."""
    linhas = [
        ("17-02-2026", "TRF P/ FULANO", -1042.69, 500.00),
        ("17-02-2026", "TRF P/ FULANO", 1042.69, 1542.69),
        ("17-02-2026", "TRF P/ FULANO", -1042.69, 500.00),
        ("16-02-2026", "Vencimento", 1542.69, 1542.69),
    ]
    previa, r = _importar(admin, "d.xls", xls_santander(linhas))
    assert "Importar 4 movimentos" in previa.text and r.status_code == 303
    # reimportar o mesmo período não duplica nada
    previa, _ = _importar(admin, "d2.xls", xls_santander(linhas))
    assert "Importar 0 movimentos" in previa.text


def xlsx_creditoagricola(linhas) -> bytes:
    """linhas: (data 'dd/mm/aaaa', descrição, valor texto, saldo texto, descritivo1), mais recente primeiro."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Conta_12345678901"
    ws.append(["Operação :", "Consulta actividade"])
    ws.append(["Número de Conta :", "12345678901"])
    ws.append([])
    ws.append([])
    ws.append(["Data Movimento", "Data Valor", "Descrição", "Valor", "Saldo após movimento", "Tipo",
               "Nome do Ordenante", "NIB/IBAN/Conta do Ordenante", "Nome do Beneficiário", "Descritivo1",
               "Descritivo2", "Referencia"])
    for d, desc, v, s, extra in linhas:
        ws.append([d, d, desc, v, s, "Débito" if v.startswith("-") else "Crédito", None, None, None, extra, extra, None])
    ws.append([])
    ws.append(["Data de impressão: 22/09/2026 | 11:23"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parser_creditoagricola():
    e = ler_extrato("Depósitos à Ordem.xlsx", xlsx_creditoagricola([
        ("31/03/2026", "Trf crédito CONDOMINIO", "-1.150,00 €", "4.352,50 €", "Condominio 1 Semestre"),
        ("30/03/2026", "DD-123-NOS Comuni", "-7,43 €", "5.502,50 €", "DD-123-NOS Comuni"),
        ("20/03/2026", "COMPRA ANDREIA &amp; ZA 7859271/25", "-22,50 €", "5.509,93 €", None),
        ("16/03/2026", "Pensão", "1.532,43 €", "5.532,43 €", None),
    ]))
    assert e.parser == "ca" and e.conta_ref == "12345678901" and not e.avisos
    assert [m.montante for m in e.movimentos] == [153243, -2250, -743, -115000]
    assert e.movimentos[1].descricao == "COMPRA ANDREIA & ZA 7859271/25"
    assert e.movimentos[2].descricao == "DD-123-NOS Comuni"                      # descritivo repetido: ignorado
    assert e.movimentos[3].descricao == "Trf crédito CONDOMINIO — Condominio 1 Semestre"
    assert e.movimentos[3].saldo == 435250
